"""
File system operations and document reading.
"""
from __future__ import annotations

import asyncio
import os
import shutil
import zipfile
from pathlib import Path
from typing import Any

from loguru import logger

try:
    import aiofiles  # type: ignore
    _AIOFILES = True
except ImportError:
    _AIOFILES = False


# ── File/folder operations ────────────────────────────────────────────────────

async def create_folder(path: str, **_: Any) -> str:
    p = Path(path)
    loop = asyncio.get_running_loop()
    await loop.run_in_executor(None, lambda: p.mkdir(parents=True, exist_ok=True))
    return f"Folder created: {p}"


async def delete_file(path: str, **_: Any) -> str:
    p = Path(path)
    loop = asyncio.get_running_loop()
    if not p.exists():
        return f"Not found: {path}"
    if p.is_dir():
        await loop.run_in_executor(None, lambda: shutil.rmtree(str(p)))
    else:
        await loop.run_in_executor(None, p.unlink)
    return f"Deleted: {path}"


async def move_file(src: str, dst: str, **_: Any) -> str:
    loop = asyncio.get_running_loop()
    await loop.run_in_executor(None, lambda: shutil.move(src, dst))
    return f"Moved {src} → {dst}"


async def copy_file(src: str, dst: str, **_: Any) -> str:
    loop = asyncio.get_running_loop()
    await loop.run_in_executor(None, lambda: shutil.copy2(src, dst))
    return f"Copied {src} → {dst}"


async def rename_file(src: str, new_name: str, **_: Any) -> str:
    p = Path(src)
    target = p.parent / new_name
    loop = asyncio.get_running_loop()
    await loop.run_in_executor(None, lambda: p.rename(target))
    return f"Renamed to {target.name}"


async def search_files(query: str, directory: str = "~", **_: Any) -> str:
    base = Path(directory).expanduser()
    loop = asyncio.get_running_loop()
    results = await loop.run_in_executor(None, _search_sync, base, query)
    if not results:
        return f"No files matching '{query}' found in {base}."
    return "Found:\n" + "\n".join(str(r) for r in results[:20])


def _search_sync(base: Path, query: str) -> list[Path]:
    results: list[Path] = []
    q = query.lower()
    try:
        for p in base.rglob("*"):
            if q in p.name.lower():
                results.append(p)
                if len(results) >= 50:
                    break
    except PermissionError:
        pass
    return results


async def zip_files(files: list[str], output: str, **_: Any) -> str:
    loop = asyncio.get_running_loop()
    await loop.run_in_executor(None, _zip_sync, files, output)
    return f"Zipped {len(files)} file(s) to {output}"


def _zip_sync(files: list[str], output: str) -> None:
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as zf:
        for f in files:
            zf.write(f, Path(f).name)


async def extract_archive(archive: str, destination: str = "", **_: Any) -> str:
    dst = destination or str(Path(archive).parent)
    loop = asyncio.get_running_loop()
    await loop.run_in_executor(None, lambda: shutil.unpack_archive(archive, dst))
    return f"Extracted to {dst}"


async def empty_recycle_bin(**_: Any) -> str:
    try:
        import winshell  # type: ignore
        winshell.recycle_bin().empty(confirm=False, show_progress=False, sound=False)
        return "Recycle bin emptied."
    except ImportError:
        import subprocess
        subprocess.run(
            ["powershell", "-Command", "Clear-RecycleBin -Force -ErrorAction SilentlyContinue"],
            check=False,
        )
        return "Recycle bin emptied."


# ── Document reading ──────────────────────────────────────────────────────────

async def read_file_content(path: str, max_chars: int = 4000, **_: Any) -> str:
    p = Path(path)
    if not p.exists():
        return f"File not found: {path}"
    suffix = p.suffix.lower()
    loop = asyncio.get_running_loop()

    if suffix == ".pdf":
        return await loop.run_in_executor(None, _read_pdf, p, max_chars)
    elif suffix in (".docx", ".doc"):
        return await loop.run_in_executor(None, _read_docx, p, max_chars)
    elif suffix in (".xlsx", ".xls"):
        return await loop.run_in_executor(None, _read_excel, p, max_chars)
    else:
        # Plain text fallback
        try:
            text = p.read_text(encoding="utf-8", errors="replace")
            return text[:max_chars]
        except Exception as exc:
            return f"Could not read file: {exc}"


def _read_pdf(path: Path, max_chars: int) -> str:
    try:
        import fitz  # PyMuPDF  # type: ignore
        doc = fitz.open(str(path))
        text = ""
        for page in doc:
            text += page.get_text()
            if len(text) >= max_chars:
                break
        return text[:max_chars]
    except ImportError:
        return "PyMuPDF not installed."
    except Exception as exc:
        return f"PDF read error: {exc}"


def _read_docx(path: Path, max_chars: int) -> str:
    try:
        from docx import Document  # type: ignore
        doc = Document(str(path))
        text = "\n".join(p.text for p in doc.paragraphs)
        return text[:max_chars]
    except ImportError:
        return "python-docx not installed."
    except Exception as exc:
        return f"DOCX read error: {exc}"


def _read_excel(path: Path, max_chars: int) -> str:
    try:
        import openpyxl  # type: ignore
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        lines: list[str] = []
        for sheet in wb.worksheets:
            lines.append(f"[Sheet: {sheet.title}]")
            for row in sheet.iter_rows(values_only=True):
                lines.append("\t".join(str(c) if c is not None else "" for c in row))
            if sum(len(l) for l in lines) >= max_chars:
                break
        return "\n".join(lines)[:max_chars]
    except ImportError:
        return "openpyxl not installed."
    except Exception as exc:
        return f"Excel read error: {exc}"
